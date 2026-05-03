from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .models import WorkbookSpec


class ExcelWorkbookBuilder:
    def build(self, spec: WorkbookSpec, output_path: Path) -> Path:
        output_path.parent.mkdir(parents=True, exist_ok=True)

        workbook = Workbook()
        summary = workbook.active
        summary.title = "Resumen"
        self._write_summary(summary, spec)

        for table_spec in spec.tables:
            sheet = workbook.create_sheet(self._safe_sheet_name(table_spec.title))
            self._write_table(sheet, table_spec)

        workbook.save(output_path)
        return output_path

    def _write_summary(self, sheet, spec: WorkbookSpec) -> None:
        sheet["A1"] = spec.title
        sheet["A1"].font = Font(size=18, bold=True, color="1F2937")
        sheet["A3"] = "Tablas"
        sheet["A3"].font = Font(bold=True)
        sheet.append([])

        sheet["A4"] = "Nombre"
        sheet["B4"] = "Filas"
        sheet["C4"] = "Columnas"
        self._style_header(sheet, 4, 3)

        for table in spec.tables:
            sheet.append([table.title, table.row_count, table.column_count])

        self._autosize(sheet)

    def _write_table(self, sheet, table_spec) -> None:
        sheet.append(table_spec.headers)
        for row in table_spec.rows:
            sheet.append(row)

        self._style_header(sheet, 1, table_spec.column_count)
        self._autosize(sheet)

        last_row = max(1, table_spec.row_count + 1)
        last_col = max(1, table_spec.column_count)
        ref = f"A1:{get_column_letter(last_col)}{last_row}"
        table = Table(displayName=self._safe_table_name(table_spec.title), ref=ref)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        sheet.add_table(table)
        sheet.freeze_panes = "A2"

    def _style_header(self, sheet, row: int, columns: int) -> None:
        fill = PatternFill("solid", fgColor="111827")
        for column in range(1, columns + 1):
            cell = sheet.cell(row=row, column=column)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center")

    def _autosize(self, sheet) -> None:
        for column_cells in sheet.columns:
            max_length = 0
            column = get_column_letter(column_cells[0].column)
            for cell in column_cells:
                if cell.value is not None:
                    max_length = max(max_length, len(str(cell.value)))
            sheet.column_dimensions[column].width = min(max(max_length + 2, 12), 42)

    def _safe_sheet_name(self, name: str) -> str:
        safe = "".join(ch for ch in name if ch not in r"[]:*?/\\").strip()
        return (safe or "Tabla")[:31]

    def _safe_table_name(self, name: str) -> str:
        safe = "".join(ch if ch.isalnum() else "_" for ch in name).strip("_")
        safe = safe or "Tabla"
        if safe[0].isdigit():
            safe = f"T_{safe}"
        return safe[:255]
